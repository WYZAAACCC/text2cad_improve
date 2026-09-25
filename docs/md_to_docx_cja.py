from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def set_run_font(run, east='宋体', latin='Times New Roman', size=10.5, bold=None, italic=None):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), east)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_inline_runs(par, text, size=10.5, east='宋体', latin='Times New Roman'):
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith('**') and part.endswith('**')
        value = part[2:-2] if bold else part
        run = par.add_run(value)
        set_run_font(run, east=east, latin=latin, size=size, bold=bold)


def add_page_number(par):
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = par.add_run()
    fld_char1 = OxmlElement('w:fldChar')
    fld_char1.set(qn('w:fldCharType'), 'begin')
    instr_text = OxmlElement('w:instrText')
    instr_text.set(qn('xml:space'), 'preserve')
    instr_text.text = ' PAGE '
    fld_char2 = OxmlElement('w:fldChar')
    fld_char2.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, size=9)


def make_document(title: str):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.1)
    sec.bottom_margin = Cm(2.1)
    sec.left_margin = Cm(2.25)
    sec.right_margin = Cm(2.25)

    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.12

    add_page_number(sec.footer.paragraphs[0])
    doc.core_properties.title = title
    doc.core_properties.subject = 'CJA paper outline and initial draft'
    return doc


def add_heading(doc, text, level):
    if level == 1:
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_after = Pt(14)
        add_inline_runs(par, text, size=18, east='黑体', latin='Arial')
        for run in par.runs:
            run.bold = True
        return
    sizes = {2: 14, 3: 12, 4: 11}
    east = '黑体'
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(9 if level == 2 else 6)
    par.paragraph_format.space_after = Pt(4)
    par.paragraph_format.keep_with_next = True
    add_inline_runs(par, text, size=sizes.get(level, 10.5), east=east, latin='Arial')
    for run in par.runs:
        run.bold = True
    return par


def add_paragraph(doc, text, style=None):
    par = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    add_inline_runs(par, text)
    return par


def add_table(doc, rows):
    if not rows:
        return
    width = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        for ci in range(width):
            value = row[ci] if ci < len(row) else ''
            cell = table.cell(ri, ci)
            cell.text = ''
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.line_spacing = 1.0
            add_inline_runs(par, value, size=8.7 if ri else 8.9)
            if ri == 0:
                for run in par.runs:
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def parse_table(lines):
    rows = []
    for line in lines:
        stripped = line.strip()
        if not stripped.startswith('|'):
            continue
        cells = [x.strip() for x in stripped.strip('|').split('|')]
        if all(re.fullmatch(r':?-{3,}:?', x) for x in cells if x):
            continue
        rows.append(cells)
    return rows


def markdown_to_docx(md_path: Path, out_path: Path):
    text = md_path.read_text(encoding='utf-8')
    title = next((ln[2:].strip() for ln in text.splitlines() if ln.startswith('# ')), md_path.stem)
    doc = make_document(title)
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        if not line.strip():
            i += 1
            continue
        if line.startswith('|'):
            end = i
            while end < len(lines) and lines[end].strip().startswith('|'):
                end += 1
            add_table(doc, parse_table(lines[i:end]))
            i = end
            continue
        if line.startswith('#### '):
            add_heading(doc, line[5:].strip(), 4)
        elif line.startswith('### '):
            add_heading(doc, line[4:].strip(), 3)
        elif line.startswith('## '):
            add_heading(doc, line[3:].strip(), 2)
        elif line.startswith('# '):
            add_heading(doc, line[2:].strip(), 1)
        elif line.startswith('> '):
            par = add_paragraph(doc, line[2:].strip())
            par.paragraph_format.left_indent = Cm(0.6)
            for run in par.runs:
                run.italic = True
                run.font.color.rgb = RGBColor(80, 80, 80)
        elif re.match(r'^\d+\.\s+', line):
            value = re.sub(r'^\d+\.\s+', '', line)
            add_paragraph(doc, value, style='List Number')
        elif line.startswith('- '):
            add_paragraph(doc, line[2:].strip(), style='List Bullet')
        elif line.strip() == '---':
            pass
        else:
            add_paragraph(doc, line.strip())
        i += 1
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    print(f'created: {out_path}')


if __name__ == '__main__':
    if len(sys.argv) != 3:
        raise SystemExit('usage: md_to_docx_cja.py input.md output.docx')
    markdown_to_docx(Path(sys.argv[1]), Path(sys.argv[2]))
