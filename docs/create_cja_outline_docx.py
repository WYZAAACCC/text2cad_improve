from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

OUT = 'docs/CJA_自动仿真与反馈修复闭环_创新点与论文大纲.docx'

doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.2)
section.bottom_margin = Cm(2.2)
section.left_margin = Cm(2.4)
section.right_margin = Cm(2.4)

style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(10.5)
style._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
style.paragraph_format.space_after = Pt(5)
style.paragraph_format.line_spacing = 1.15


def set_run_font(run, east='宋体', latin='Times New Roman', size=10.5, bold=None):
    run.font.name = latin
    run._element.rPr.rFonts.set(qn('w:eastAsia'), east)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def p(text='', bold=False, size=10.5, align=None, color=None, italic=False):
    para = doc.add_paragraph()
    if align is not None:
        para.alignment = align
    run = para.add_run(text)
    set_run_font(run, size=size, bold=bold)
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)
    return para


def h1(text):
    para = doc.add_heading(level=1)
    run = para.add_run(text)
    set_run_font(run, east='黑体', latin='Arial', size=15, bold=True)
    return para


def h2(text):
    para = doc.add_heading(level=2)
    run = para.add_run(text)
    set_run_font(run, east='黑体', latin='Arial', size=12.5, bold=True)
    return para


def h3(text):
    para = doc.add_heading(level=3)
    run = para.add_run(text)
    set_run_font(run, east='黑体', latin='Arial', size=11, bold=True)
    return para


def bullet(text, level=0):
    para = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    run = para.add_run(text)
    set_run_font(run)
    return para


def numbered(text):
    para = doc.add_paragraph(style='List Number')
    run = para.add_run(text)
    set_run_font(run)
    return para


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, text in enumerate(headers):
        hdr[i].text = ''
        run = hdr[i].paragraphs[0].add_run(text)
        set_run_font(run, east='黑体', latin='Arial', size=9.5, bold=True)
    for row in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = ''
            run = cells[i].paragraphs[0].add_run(str(text))
            set_run_font(run, size=9.2)
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    return t

# Cover
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run('CJA 论文创新点与写作大纲')
set_run_font(r, east='黑体', latin='Arial', size=20, bold=True)
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run('面向航空发动机涡轮盘的自动仿真—反馈—修复闭环')
set_run_font(r, east='黑体', latin='Arial', size=14, bold=True)
p('日期：2026-09-20', align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
p('目标期刊：Chinese Journal of Aeronautics (CJA)', align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
p('定位：生成部分仅作背景；核心贡献放在 CAD–FEA 自动仿真、反馈归因、特征修复、预测判伪与工程经验积累。',
  align=WD_ALIGN_PARAGRAPH.CENTER, italic=True)

h1('0. 执行结论')
p('现有工作已经覆盖“LLM 自动生成 FEA/CAE 脚本、调用求解器、读取结果并迭代”的一般流程，因此“用 LLM 做自动仿真”本身不再足以作为主创新。你的工作真正有差异度的部分不是多智能体数量，也不是自然语言生成输入文件，而是以下三层证据链：')
numbered('设计证据：feature-level CAD dependency graph，保证修改的是正确的涡轮盘特征或参数。')
numbered('几何证据：pre-solve STEP solid-effect gate，确认修改确实进入最终实体目标区域，而不是只改了文档字段。')
numbered('物理证据：下一轮独立 ANSYS 求解 + mesh noise floor + prediction scoring，判定 confirmed/refuted/unmoved。')
p('建议把论文主叙事从“自动仿真 agent”调整为“面向涡轮盘的、可判伪的 CAD–FEA 反馈修复框架”。最突出的创新点是：feature-aware repair + solid-effect verification + falsifiable engineering experience learning。')

h1('1. 已有工作版图')
h2('1.1 最接近的工作')
table(
    ['工作', '年份/期刊', '主要贡献', '它已经覆盖了什么', '你的差异空间'],
    [
        ['CFD-copilot', '2026, CJA', '领域适配 LLM + MCP，增强 CFD 仿真自动化', '自然语言到 CFD 工作流、工具编排、仿真自动化', '你聚焦结构 FEA 和涡轮盘；不是替代 CFD-copilot，而是把闭环推进到 CAD feature 修复、实体效果验证和预测判伪'],
        ['AbaqusAgent', '2026, arXiv/多 agent', '六个 agent 完成 Abaqus 输入、运行、审查、可视化', '通用固体力学 FEA case 生成和执行，50 例，86% 成功', '你没有停留在一次性生成；重点是跨 revision 修复、最终实体验证和真实设计经验积累'],
        ['ALL-FEM', '2026, CMAME', '微调 LLM + agentic FEM 代码生成', 'FEniCS、多物理、多 agent、运行反馈', '你的领域不再是通用代码生成，而是涡轮盘 feature edit、载荷路径、网格/面映射和数据归因'],
        ['AutoFEA', '2025, AAAI', 'LLM + GCN-Transformer 检索，生成 CalculiX 输入', '降低 FEA 代码生成幻觉、提高输入文件正确率', '你关注“修什么设计变量、是否改到真实几何、下一次物理结果是否支持预测”，而不是只生成正确输入文件'],
        ['PAMF', '2026, JMST', '微调 LLM agent 预测网格误差并自适应加密', 'a priori mesh generation、error feedback、APDL 代码库', '你覆盖完整 CAD–FEA–feedback–repair 循环，并且有构件级应力结果和设计变量修改'],
        ['LLM Mechanical Designer', '2024/2026', 'LLM + FEM 迭代优化 2D truss', '设计生成、FEM 反馈、结构修改、与 NSGA-II 比较', '你的测试对象是真实涡轮盘 CAD/FEA，加入 feature dependency 和工程经验规则；不要只重复“LLM 优化结构”'],
        ['Next-generation CAE', '2025, CMAME', '综述 LLM 驱动的下一代 CAE', '框架、机会和挑战', '你提供特定航空航天部件上的可运行闭环和真实证据，而不是综述'],
        ['Generic CAE harness', '2026, arXiv', '比较通用 harness、多 agent、执行反馈和领域知识', '执行反馈有效；脚本化 reflection 可能无效；领域教程有帮助', '你的贡献应是“哪些领域特定的 measurement/gate 不能被通用 harness 替代”，而不是“增加 agent 数量”'],
    ],
    widths=[2.6,2.2,3.7,4.3,4.7]
)

h2('1.2 涡轮盘结构设计已有工作')
table(
    ['方向', '代表性工作', '已有能力', '你的论文不应重复的主张'],
    [
        ['应力约束拓扑优化', 'CJA 2023: turbine disk maximum stress prediction and constraints', '拓扑优化中的最大应力预测和约束处理', '不要声称首次做涡轮盘应力优化'],
        ['通风孔/开口形状优化', 'CJA 2021: elliptical ventilation openings; AST 2019: non-circular vent hole', '解析/FEM/代理模型优化开孔位置和形状', '不要声称首次优化孔和应力集中'],
        ['双腹板盘优化', 'SMO 2022: topology and shape optimization of twin-web turbine disk', '形状、拓扑、Kriging、多学科优化', '不要声称首次用代理模型做盘优化'],
        ['疲劳/蠕变可靠性与寿命', 'IJF 2021; RESS 2022: fatigue/creep-fatigue reliability of turbine disks', '可靠性、寿命、概率评估、热点识别', '你的系统目前仍以应力和位移为主，不要把结果包装成寿命优化已完成'],
        ['生成式概念设计', 'CJA 2025: generative AI aircraft conceptual design', 'LLM 生成概念方案并与工程师对比', '你的重点不是概念设计，而是求解后的定量反馈和修复'],
    ],
    widths=[3.0,5.0,4.5,5.0]
)

h1('2. 建议的 CJA 创新点')
h2('C1. 面向涡轮盘的特征级 CAD 参数依赖图与联合修复')
p('现有 LLM/CAE agent 多数把设计修改视为“改脚本”或“改一个标量”。涡轮盘的几何语义不是标量：meridian mirror vertex、slot profile、fillet family、hole pattern、rim thickness 等之间具有设计依赖。你当前已经有：')
bullet('mirror_vertex_pair：同一回转面的镜像顶点必须成对移动；')
bullet('fillet_family：同一 feature 的圆角族可以联合修改；')
bullet('pattern_instances：孔/槽 circular pattern 的 source 和 placement 关系；')
bullet('profile_contour：同一闭合轮廓的坐标关系；')
bullet('feature_bundle：一个生成 feature 的可编辑参数范围。')
p('论文贡献可以表述为：提出一种从 CAD document 自动构造 feature-level dependency graph，并允许 agent 在同一 finding 下执行多参数、原子式、可审计修复的方法。它不是“更大 prompt”，而是把工程几何语义显式提供给 agent。')

h2('C2. 求解前真实实体效果门：Document edit ≠ Solid change')
p('这是你目前最有辨识度的工程贡献之一。很多 agent 只能证明自己写入了参数，不能证明最终 CAD/STEP 真的变了。你加入了：')
bullet('最终 STEP 的 boundary triangle cells；')
bullet('目标半径/轴向窗口内的局部 surface area 和 radial boundary 比较；')
bullet('target window 与 finding 的 radius/z/parameter reach 绑定；')
bullet('如果目标区域未变化，则在 ANSYS solve 前 fail-closed。')
p('可以把这概括为：“把设计修复的验证从文本/文档层提升到最终 solid feature 层”。在 CJA 语境下，这比单纯增加 agent 数量更有工程价值。')

h2('C3. 可判伪的反馈—修复协议')
p('现有 CAE agent 通常说“读取结果并迭代”，但很少定义反馈何时应被相信。你的协议是：')
numbered('feedback agent 给出 mechanism、location、change 和 prediction；')
numbered('evidence 必须能在真实结果场中重读；')
numbered('revise agent 只能写 document 中真实存在的参数；')
numbered('写入后必须产生真实 diff，并且数值/相对变化/镜像关系一致；')
numbered('下一 revision 用独立 solve 检查 prediction；')
numbered('用 mesh noise floor 区分 confirmed、partial、refuted、unmoved。')
p('这形成一个“physical prediction → intervention → re-solve → refutation/confirmation”的闭环。论文里可以把它称为 physics-grounded falsifiable feedback。')

h2('C4. 面向航空航天工程的经验学习系统')
p('建议把“engineering experience learning”写成核心候选创新，而不是附属功能。它的差异点不是普通 memory，而是：')
bullet('rule 只能来自真实 FEA run；')
bullet('rule 按 mechanism、parameter、direction 存储；')
bullet('规则状态从 seed → candidate → trusted/refuted；')
bullet('refutation 必须有真实 observation 和 side effect；')
bullet('跨族 transfer 后才能称为泛化；')
bullet('规则带有网格噪声底、载荷完整性和几何可达性约束。')
p('这已经区别于 ExpeL、Reflexion 这类通用文本经验学习：这里的 experience 是带物理量纲、误差底和设计参数的工程证据。')

h2('C5. 多设计族机场/涡轮盘 benchmark')
p('当前工作已经有 11 族选面统计、32 族 dependency graph 覆盖实验和真实 D19/D27 solve。建议把它包装为：')
bullet('frozen face-selection benchmark；')
bullet('feature dependency coverage benchmark；')
bullet('closed-loop prediction benchmark（confirmed/partial/refuted/unmoved）；')
bullet('knowledge transfer benchmark（seed rule 从一个族转移到另一个族）。')
p('这比只展示一个 D27 case 更像 CJA 论文，也能把负结果写成贡献：D19 bore-opening prediction 被真实实验推翻，并发现 load-surface side effect。')

h1('3. 建议的论文标题')
bullet('Feature-Aware and Falsifiable CAD–FEA Closed Loop for Aero-Engine Turbine Disk Design Repair')
bullet('From Simulation to Design Repair: An Evidence-Grounded Agentic Framework for Turbine Disk Stress-Field Feedback')
bullet('Engineering Experience Learning through Falsifiable CAD–FEA Feedback for Aero-Engine Turbine Disks')
bullet('A Feature-Dependency and Prediction-Scoring Agent Framework for Turbine Disk Structural Simulation and Repair')
p('推荐第一或第二个标题。标题里出现 feature-aware、falsifiable、closed-loop、turbine disk，比“LLM agent”更能体现 CJA 关心的工程和可信度。')

h1('4. CJA 论文大纲（英文写作结构 + 中文说明）')
h2('Abstract')
p('推荐结构：背景一句；已有 LLM/CAE 工作一句；指出 attribution gap：参数修改、最终 solid、物理响应之间没有闭合；提出 feature-aware/falsifiable loop；给出 3 个定量结果；最后一句给工程意义。')
p('必须出现的数字：feature dependency coverage；face-selection exact accuracy；real closed-loop solve；prediction outcomes；revise patch correctness；至少一个 refuted case。不要只写 qualitative claims。')

h2('1. Introduction')
bullet('1.1 Aero-engine turbine disk design and the cost of repeated CAD–FEA iterations.')
bullet('1.2 LLM/agentic automation has advanced CAE workflow generation.')
bullet('1.3 Remaining attribution gap: document edit ≠ solid change ≠ physical improvement.')
bullet('1.4 This paper: feature-level dependency graph, pre-solve solid-effect gate, falsifiable feedback, engineering experience memory.')
bullet('1.5 Contributions and paper organization.')
p('引言故事建议：不要从“LLM 很热”开始，而从涡轮盘的实际工程链条开始：face selection → load path → mesh → stress peak → feature edit → re-solve → rule reuse。')

h2('2. Related Work and Positioning')
bullet('2.1 Turbine disk structural optimization: topology, shape, surrogate, reliability, fatigue/creep.')
bullet('2.2 LLM/agentic CAE and FEA: CFD-copilot, AbaqusAgent, ALL-FEM, AutoFEA, PAMF, OpenFOAMGPT.')
bullet('2.3 Iterative design and experience learning: LLM agent as mechanical designer, ExpeL, Reflexion, CRITIC, case-based engineering knowledge reuse.')
bullet('2.4 Positioning: what current systems automate vs what this paper verifies and repairs.')
p('结尾放一张 Position Matrix 表：Feature-aware / Solid-effect verified / Falsifiable prediction / Noise-aware / Cross-family knowledge / Turbine-disk domain。')

h2('3. Problem Formulation')
p('定义 document state、simulation experiment、finding、change、prediction、outcome。建议形式化：')
bullet('D_N: CAD document at revision N；')
bullet('E: simulation experiment held fixed across revisions；')
bullet('M_N = S(D_N, E): physical metrics;')
bullet('F_N: finding and prediction about intervention ΔD;')
bullet('G(D_N, ΔD): geometric validity predicate;')
bullet('O_N ∈ {confirmed, partial, refuted, unmoved, unmeasurable, untested}。')
p('明确指出：论文的目标不是证明某个规则永远正确，而是给出 evidence-grounded intervention and refutation protocol。')

h2('4. Methodology')
h3('4.1 Overall closed-loop architecture')
p('generate → frame/domain/setup/facefind/mesh → solve → postprocess/verify → feedback → revise → next generate。生成部分只写一段，强调它是已有的 document generator。')
h3('4.2 Frame normalization and simulation consistency')
p('profile/Gmsh/mesh mapping/APDL/postprocess 使用同一个 case frame；说明非全局 Z 的契约、fail-closed 和当前 topology 限制。')
h3('4.3 Measurement-driven simulation agents')
p('face selection：flank normal family、outer radial band、radial load alignment；mesh agent：refinement、convergence/noise floor；postprocess：load audit、radial bands、stress components。')
h3('4.4 Feature-level CAD dependency graph')
p('定义 nodes/edges/parameter groups；说明 mirror、fillet family、pattern、contour、feature bundle；给算法伪代码。')
h3('4.5 Feedback agent')
p('region ranking、mechanism signature、evidence re-read、change proposal、prediction。')
h3('4.6 Revise agent')
p('document parameter resolution、absolute/relative magnitude authority、mirror pair、multi-parameter joint edit、no-op refusal、master isolation。')
h3('4.7 Pre-solve solid-effect verification')
p('triangle cells、target window、changed cells、radial boundary、fail-closed。这里是 engineering contribution，要画流程图。')
h3('4.8 Falsifiable prediction scoring and engineering experience memory')
p('mesh noise floor、confirmed/partial/refuted/unmoved、rule lifecycle、evidence count、cross-family generalization。')

h2('5. Experimental Setup')
bullet('5.1 Dataset and design families: D19/D27 verified; D15–D29 extended families; D01–D32 dependency coverage.')
bullet('5.2 Hardware/software: ANSYS 18.1, Gmsh, Flash model, Python harness.')
bullet('5.3 Metrics: exact-set accuracy, wrong accepted, Wilson CI, submission rate, calls; feedback evidence rate; revise patch accuracy; prediction outcome distribution; rule sample count.')
bullet('5.4 Baselines: one-shot LLM, execution-only repair, generic harness, multi-agent without feature graph, surrogate/Kriging or Bayesian optimization for selected cases.')

h2('6. Results')
h3('6.1 Feature dependency coverage across 32 families')
p('报告 mirror/fillet/contour/feature 32/32；pattern 28/32（D01–D04 无 pattern）；group count median 19。')
h3('6.2 Face-selection benchmark')
p('报告 51/51 exact、0 wrong accepted、Wilson interval；明确 D19/D27 verified vs provisional references。')
h3('6.3 Pre-solve solid-effect verification')
p('报告 D27 bore 60→63 的 target min radius、area、changed cells；说明 no-op patch 不能进入 solve。')
h3('6.4 Closed-loop turbine disk case studies')
p('D27：bore 60→66，peak -1.22%，但 noise floor 3.81%，判定 unmoved。D19：60→63→66，peak 1149.708→1150.936→1152.771 MPa，两次 prediction refuted，load-surface stress +6.09%/+5.06%。')
h3('6.5 Feedback and revise reliability')
p('feedback 7/7 accepted、evidence 58/58；revise applied patch 3/3 grounded；document patch correctness 10/11；master changed 0。')
h3('6.6 Engineering experience learning')
p('展示 seed→candidate→refuted 的状态变化；当前 1 条真实 refuted rule；讨论 trusted 所需样本数。')
h3('6.7 Ablation study')
p('必须做：no feature graph、no solid-effect gate、no prediction scoring、no noise floor、single-agent vs multi-agent、generic harness vs domain-measured harness。')

h2('7. Discussion')
bullet('Why feature-aware measurement is necessary in aero-engine CAD–FEA.')
bullet('Why execution success is not physical success.')
bullet('Why refutation and unmoved outcomes are scientifically useful.')
bullet('When an agent should needs_input or skip rather than invent a design change.')
bullet('Engineering implications for turbine disk design reuse and certification.')

h2('8. Limitations and Future Work')
bullet('Current real closed-loop evidence is concentrated in D27/D19; need more families.')
bullet('Trusted rules are not yet statistically mature.')
bullet('Non-global-Z solver-level validation is blocked by rigid-transform topology support.')
bullet('Current objective is primarily stress/displacement; extend to fatigue, creep, burst margin, life and uncertainty.')
bullet('Future: multi-objective Pareto repair, uncertainty-aware feedback, expert-in-the-loop validation.')

h2('9. Conclusion')
p('用三句话收束：framework、evidence、engineering significance。不要重复摘要。')

h1('5. 投稿前必须补的实验（按优先级）')
h2('P0：支撑主创新')
numbered('至少再跑 4–6 个涡轮盘设计族的完整两 revision 闭环；每个 case 固定 experiment，并报告 noise floor。')
numbered('形成至少 2 条 trusted、2 条 refuted、2 条跨族复用规则；每条规则列出真实 run、方向、参数、样本数和 side effects。')
numbered('完成 ablation：去掉 feature graph、去掉 solid-effect gate、去掉 prediction/noise floor，各自对 exact patch、wrong attribution、refuted detection 的影响。')
numbered('增加 baseline：one-shot LLM、execution-only repair、generic harness、Kriging/Bayesian optimization 或人工工程师。')
numbered('增加工程验证：载荷守恒、mesh independence、关键案例人工复核、制造/几何可行性拒绝率。')
h2('P1：增强 CJA 说服力')
numbered('扩展 D19/D27 的 fatigue/creep 或 safety margin 指标，至少作为 discussion。')
numbered('给出跨族 negative transfer 分析：什么规则从 D27 到 D19 会失败。')
numbered('把 D19 bore-opening 失败 case 作为 main case study，而不是藏在附录。')
numbered('非 Z 若无法完成 OCAF transform，就在 limitations 明确说明，不要把它写成已验证的工业能力。')

h1('6. 审稿人最可能的质疑与回应')
table(
    ['审稿人质疑', '建议回应'],
    [
        ['“LLM/CAE automation 已经很多人做了。”', '不要把贡献写成首个自动 FEA；写成首个面向涡轮盘 feature 修复的 solid-effect + falsifiable loop。'],
        ['“这是不是只是 prompt engineering？”', '强调 harness 中的 measurement、dependency graph、solid-effect gate、noise-floor scoring 和 fail-closed protocol。'],
        ['“为什么需要 agent，不用贝叶斯优化？”', 'Agent 处理的是非结构化诊断、特征语义、不可达参数和模型修复；BM/Bayesian 作为对照实验证明边界。'],
        ['“为什么结果有时 refuted 或 unmoved？”', '这正是贡献：系统能识别不真实改进；refutation 和 unmoved 是物理证据，不是失败。'],
        ['“知识学习是否只是 memory？”', '强调规则必须来自真实 run，带有 mechanism/parameter/direction、noise、side effects 和跨族 transfer。'],
        ['“涡轮盘领域太窄。”', '窄领域正适合 CJA；先把涡轮盘的可信闭环做深，再讨论通用化。'],
    ],
    widths=[5.0,11.5]
)

h1('7. 建议写入摘要的定量结果')
bullet('Feature dependency coverage: 32 families; mirror/fillet/contour/feature 100%; pattern 28/32.')
bullet('Face selection: 51/51 exact; 0 wrong accepted; submission 100%; Wilson 95% CI [0.930, 1.0].')
bullet('Pre-solve solid effect: D27 bore 60→63; target min radius 59.856→62.819 mm; local area 199259.7→198374.9 mm2; target changed.')
bullet('Closed-loop D19: 1149.708→1150.936→1152.771 MPa; two bore-opening predictions refuted; load-surface stress side effects +6.09%/+5.06%.')
bullet('D27 paired experiment: 1709.359→1688.457 MPa (-1.22%), but noise floor 3.81%, verdict unmoved.')
bullet('Feedback/revise: 7/7 accepted, evidence 58/58; applied patches 3/3 grounded, master changed 0.')
bullet('Prediction outcomes: 2 confirmed, 4 refuted, 1 unmoved; 1 real refuted rule with sample.')
p('注意：摘要里不要写“successful automatic optimization”，应写“evidence-grounded feedback and repair with confirmation/refutation”。')

h1('8. 参考文献候选（用于 Related Work 起步）')
refs = [
'[1] CFD-copilot: Leveraging domain-adapted large language model and model context protocol to enhance simulation automation. Chinese Journal of Aeronautics, 2026. DOI: 10.1016/j.cja.2026.104321.',
'[2] AbaqusAgent: A Multi-AI-agent Framework Enabling End-to-end Finite Element Analysis for Solid Mechanics Problems. arXiv:2606.00138, 2026.',
'[3] ALL-FEM: Agentic Large Language Models Fine-tuned for Finite Element Methods. Computer Methods in Applied Mechanics and Engineering, 2026. DOI: 10.1016/j.cma.2026.118985.',
'[4] AutoFEA: Enhancing AI Copilot by Integrating Finite Element Analysis Using Large Language Models with Graph Neural Networks. AAAI 2025. DOI: 10.1609/aaai.v39i22.34582.',
'[5] PAMF: An LLM-driven framework for automated mesh generation in mechanical simulation and CAE workflows. Journal of Mechanical Science and Technology, 2026. DOI: 10.1007/s12206-026-0334-6.',
'[6] MechAgents: Large language model multi-agent collaborations can solve mechanics problems, generate new data, and integrate knowledge. Extreme Mechanics Letters, 2024. DOI: 10.1016/j.eml.2024.102131.',
'[7] Large Language Model Agent as a Mechanical Designer. Journal of Engineering Design, 2026 / arXiv:2404.17525. DOI: 10.1080/09544828.2026.2624356.',
'[8] Large language model-empowered next-generation computer-aided engineering. Computer Methods in Applied Mechanics and Engineering, 2025. DOI: 10.1016/j.cma.2025.118591.',
'[9] What Do CAE Simulation Agents Really Need Beyond a Generic Harness? arXiv:2609.03718, 2026.',
'[10] OpenFOAMGPT: A retrieval-augmented large language model (LLM) agent for OpenFOAM-based computational fluid dynamics. Physics of Fluids, 2025. DOI: 10.1063/5.0257555.',
'[11] A rapidly structured aircraft concept design method based on generative artificial intelligence. Chinese Journal of Aeronautics, 2025. DOI: 10.1016/j.cja.2025.103629.',
'[12] Topology optimization of turbine disk considering maximum stress prediction and constraints. Chinese Journal of Aeronautics, 2023. DOI: 10.1016/j.cja.2023.03.019.',
'[13] Optimal location and shape definition of elliptical ventilation openings on aero engine turbine rotors with stress concentration effect. Chinese Journal of Aeronautics, 2021. DOI: 10.1016/j.cja.2021.05.003.',
'[14] Surrogate-based optimization with improved support vector regression for non-circular vent hole on aero-engine turbine disk. Aerospace Science and Technology, 2019. DOI: 10.1016/j.ast.2019.105332.',
'[15] Topology and shape optimization of twin-web turbine disk. Structural and Multidisciplinary Optimization, 2022. DOI: 10.1007/s00158-021-03147-z.',
'[16] A unified fatigue reliability-based design optimization framework for aircraft turbine disk. International Journal of Fatigue, 2021. DOI: 10.1016/j.ijfatigue.2021.106422.',
'[17] A data-driven roadmap for creep-fatigue reliability assessment and its implementation in low-pressure turbine disk at elevated temperatures. Reliability Engineering & System Safety, 2022. DOI: 10.1016/j.ress.2022.108523.',
'[18] ExpeL: LLM Agents Are Experiential Learners. AAAI 2024. DOI: 10.1609/aaai.v38i17.29936.',
'[19] Reflexion: Language Agents with Verbal Reinforcement Learning. arXiv:2303.11366, 2023.',
'[20] Self-Refine: Iterative Refinement with Self-Feedback. arXiv:2303.17651, 2023.',
'[21] CRITIC: Large Language Models Can Self-Correct with Tool-Interactive Critiquing. arXiv:2305.11738, 2023.',
'[22] A study in applying case-based reasoning to engineering design: Mechanical bearing design. Artificial Intelligence for Engineering Design, Analysis and Manufacturing, 2003. DOI: 10.1017/s0890060403173064.',
'[23] An engineering design knowledge reuse methodology using process modelling. Research in Engineering Design, 2007. DOI: 10.1007/s00163-007-0028-8.',
'[24] Multidisciplinary Design Optimization of Turbine Disks Based on ANSYS Workbench Platforms. Procedia Engineering, 2015. DOI: 10.1016/j.proeng.2014.12.659.',
'[25] An LLM-guided structural optimization method for aero-engine lubrication hydrocyclones. Separation and Purification Technology, 2026. DOI: 10.1016/j.seppur.2026.137096.',
]
for ref in refs:
    p(ref)

h1('附录：检索方法与证据边界')
bullet('检索源：OpenAlex API、Crossref API、arXiv API、Elsevier article metadata；检索日期：2026-09-20。')
bullet('检索主题：LLM/CAE、LLM/FEA、agentic simulation、automatic boundary conditions、turbine disk optimization、surrogate/optimization、engineering knowledge reuse、experience learning。')
bullet('部分 Elsevier 论文摘要未公开返回；对这些工作使用标题、期刊元数据和可获取摘要定位，不把标题推断成具体算法细节。')
bullet('“已有工作”结论用于定位和写作策略，不替代投稿前的逐篇全文精读；正式 Related Work 前应下载全文并核验方法、数据集和指标。')
bullet('你的工作尚未达到“稳定自动应力优化器”的成熟度，因此论文应定位为 framework + evidence + case study，而不是工业级 optimizer。')

doc.save(OUT)
print(OUT)
